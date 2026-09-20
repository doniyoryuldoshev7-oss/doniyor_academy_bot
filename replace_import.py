from pathlib import Path

p = Path("app/handlers/admin.py")
s = p.read_text(encoding="utf-8")

start_marker = "# -------------------- EXCEL / CSV IMPORT --------------------"
end_marker = "# -------------------- CATALOG / STATS --------------------"

start = s.index(start_marker)
end = s.index(end_marker)

new_import = r'''# -------------------- EXCEL / CSV IMPORT --------------------
@router.callback_query(F.data == "adm:import")
async def import_start(c: CallbackQuery, state: FSMContext):
    if not is_admin(c.from_user.id):
        return

    await state.set_state(AdminState.import_file)

    await c.message.edit_text(
        "<b>Excel/CSV orqali test import</b>\n\n"
        "Faylni yuboring: <b>.xlsx</b> yoki <b>.csv</b>\n\n"
        "<b>Ustunlar:</b>\n"
        "<code>Fan | Mavzu | Savol | Savol rasmi | A | B | C | D | "
        "Savol turi | To'g'ri javob | To'g'ri javob matni | Izoh</code>\n\n"
        "Eski oddiy format ham ishlaydi.\n"
        "Savol turi: <code>closed</code> yoki <code>open</code>.\n"
        "Savol matni yoki Savol rasmi bo'lishi kerak.\n"
        "Open savolda A/B/C/D shart emas.",
        reply_markup=cancel_kb(),
    )
    await c.answer()


@router.message(AdminState.import_file, F.document)
async def import_document(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    doc = m.document
    filename = (doc.file_name or "").lower()

    if not (filename.endswith(".xlsx") or filename.endswith(".csv")):
        await m.answer(
            "Faqat .xlsx yoki .csv fayl yuboring.",
            reply_markup=cancel_kb(),
        )
        return

    import tempfile
    from pathlib import Path

    path = (
        Path(tempfile.gettempdir())
        / f"doniyor_import_{m.from_user.id}_"
          f"{doc.file_unique_id}{Path(filename).suffix}"
    )

    try:
        await m.bot.download(doc, destination=path)

        if filename.endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )
            rows = list(
                wb.active.iter_rows(values_only=True)
            )
            wb.close()

        else:
            import csv
            import io

            raw = path.read_bytes()
            csv_text = None

            for enc in (
                "utf-8-sig",
                "utf-8",
                "cp1251",
            ):
                try:
                    csv_text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    pass

            if csv_text is None:
                raise ValueError(
                    "CSV kodlashini o'qib bo'lmadi."
                )

            rows = list(
                csv.reader(io.StringIO(csv_text))
            )

        if len(rows) < 2:
            raise ValueError(
                "Faylda savollar mavjud emas."
            )

        def normalize_header(value):
            if value is None:
                return ""

            value = str(value).strip().lower()

            for code in (
                0x2018,
                0x2019,
                0x02BB,
                0x02BC,
                0x0060,
            ):
                value = value.replace(
                    chr(code),
                    "'",
                )

            value = value.replace(
                chr(0xFEFF),
                "",
            )

            return " ".join(value.split())

        headers = [
            normalize_header(x)
            for x in rows[0]
        ]

        aliases = {
            "fan": {
                "fan",
                "subject",
            },
            "mavzu": {
                "mavzu",
                "topic",
            },
            "savol": {
                "savol",
                "question",
                "text",
            },
            "image": {
                "savol rasmi",
                "savol rasmi yo'li",
                "savol rasmi yoli",
                "image",
                "image path",
                "image_path",
                "rasm",
                "rasm yo'li",
                "rasm yoli",
            },
            "a": {
                "a",
                "variant a",
            },
            "b": {
                "b",
                "variant b",
            },
            "c": {
                "c",
                "variant c",
            },
            "d": {
                "d",
                "variant d",
            },
            "mode": {
                "savol turi",
                "question type",
                "question_type",
                "type",
                "mode",
            },
            "correct": {
                "to'g'ri javob",
                "tog'ri javob",
                "togri javob",
                "correct",
                "correct option",
                "correct answer",
                "javob",
                "to'g'ri",
                "togri",
            },
            "correct_answer": {
                "to'g'ri javob matni",
                "tog'ri javob matni",
                "togri javob matni",
                "correct answer text",
                "correct_answer",
                "answer text",
                "javob matni",
            },
            "izoh": {
                "izoh",
                "explanation",
                "comment",
            },
        }

        idx = {}

        for key, names in aliases.items():
            normalized_names = {
                normalize_header(x)
                for x in names
            }

            for i, header in enumerate(headers):
                if header in normalized_names:
                    idx[key] = i
                    break

        required = [
            "fan",
            "mavzu",
        ]

        missing = [
            x for x in required
            if x not in idx
        ]

        if missing:
            raise ValueError(
                "Majburiy ustunlar topilmadi: "
                + ", ".join(missing)
            )

        def cell(row, key):
            i = idx.get(key)

            if i is None:
                return ""

            if i >= len(row):
                return ""

            if row[i] is None:
                return ""

            return str(row[i]).strip()

        valid = []
        errors = []

        for row_no, row in enumerate(
            rows[1:],
            start=2,
        ):
            if not any(
                x is not None
                and str(x).strip()
                for x in row
            ):
                continue

            subject_name = cell(row, "fan")
            topic_name = cell(row, "mavzu")
            question = cell(row, "savol")
            image_value = cell(row, "image")

            option_a = cell(row, "a")
            option_b = cell(row, "b")
            option_c = cell(row, "c")
            option_d = cell(row, "d")

            mode = cell(row, "mode").lower()
            correct = cell(row, "correct").upper()
            correct_answer = cell(
                row,
                "correct_answer",
            )
            explanation = (
                cell(row, "izoh")
                or None
            )

            if not mode:
                mode = "closed"

            if mode in {
                "yopiq",
                "variantli",
                "test",
            }:
                mode = "closed"

            if mode in {
                "ochiq",
                "free",
                "free text",
            }:
                mode = "open"

            if mode not in {
                "closed",
                "open",
            }:
                errors.append(
                    f"{row_no}-qator: "
                    "Savol turi closed yoki open "
                    "bo'lishi kerak."
                )
                continue

            if not subject_name:
                errors.append(
                    f"{row_no}-qator: Fan bo'sh."
                )
                continue

            if not topic_name:
                errors.append(
                    f"{row_no}-qator: Mavzu bo'sh."
                )
                continue

            if not question and not image_value:
                errors.append(
                    f"{row_no}-qator: Savol yoki "
                    "Savol rasmi bo'lishi kerak."
                )
                continue

            if mode == "closed":
                if not all(
                    [
                        option_a,
                        option_b,
                        option_c,
                        option_d,
                    ]
                ):
                    errors.append(
                        f"{row_no}-qator: closed savolda "
                        "A/B/C/D to'liq bo'lishi kerak."
                    )
                    continue

                if correct not in {
                    "A",
                    "B",
                    "C",
                    "D",
                }:
                    errors.append(
                        f"{row_no}-qator: To'g'ri javob "
                        "A/B/C/D bo'lishi kerak."
                    )
                    continue

                correct_answer = None

            else:
                if not correct_answer:
                    errors.append(
                        f"{row_no}-qator: open savolda "
                        "To'g'ri javob matni bo'lishi kerak."
                    )
                    continue

                correct = None

            valid.append(
                {
                    "subject": subject_name,
                    "topic": topic_name,
                    "question": question or None,
                    "image": image_value or None,
                    "mode": mode,
                    "a": option_a or None,
                    "b": option_b or None,
                    "c": option_c or None,
                    "d": option_d or None,
                    "correct": correct,
                    "correct_answer": correct_answer,
                    "explanation": explanation,
                }
            )

        if errors:
            preview = "\n".join(
                errors[:10]
            )

            more = ""

            if len(errors) > 10:
                more = (
                    f"\n... yana "
                    f"{len(errors) - 10} ta xato"
                )

            await m.answer(
                "<b>Faylda xato bor.</b>\n\n"
                f"{preview}{more}\n\n"
                "Import bajarilmadi. "
                "Faylni tuzatib qayta yuboring.",
                reply_markup=cancel_kb(),
            )
            return

        if not valid:
            raise ValueError(
                "Import qilinadigan savol topilmadi."
            )

        async with SessionLocal() as s:
            subject_cache = {}
            topic_cache = {}

            for item in valid:
                subject_name = item["subject"]
                topic_name = item["topic"]

                skey = subject_name.casefold()

                subject = subject_cache.get(
                    skey
                )

                if not subject:
                    subject = await s.scalar(
                        select(Subject).where(
                            func.lower(
                                Subject.name
                            )
                            == subject_name.lower()
                        )
                    )

                    if not subject:
                        subject = Subject(
                            name=subject_name
                        )
                        s.add(subject)
                        await s.flush()

                    subject_cache[skey] = subject

                tkey = (
                    subject.id,
                    topic_name.casefold(),
                )

                topic = topic_cache.get(tkey)

                if not topic:
                    topic = await s.scalar(
                        select(Topic).where(
                            Topic.subject_id
                            == subject.id,
                            func.lower(
                                Topic.name
                            )
                            == topic_name.lower(),
                        )
                    )

                    if not topic:
                        topic = Topic(
                            subject_id=subject.id,
                            name=topic_name,
                        )
                        s.add(topic)
                        await s.flush()

                    topic_cache[tkey] = topic

                image_path_value = None

                if item["image"]:
                    image_path_value = (
                        Path(
                            item["image"]
                        ).name
                    )

                    image_path_value = (
                        Path(
                            "app/question_images"
                        )
                        / image_path_value
                    ).as_posix()

                s.add(
                    Question(
                        topic_id=topic.id,
                        text=item["question"],
                        image_path=image_path_value,
                        question_mode=item["mode"],
                        option_a=item["a"],
                        option_b=item["b"],
                        option_c=item["c"],
                        option_d=item["d"],
                        correct_option=item["correct"],
                        correct_answer=item[
                            "correct_answer"
                        ],
                        explanation=item[
                            "explanation"
                        ],
                    )
                )

            await s.commit()

        await state.clear()

        await m.answer(
            "<b>IMPORT MUVAFFAQIYATLI!</b>\n\n"
            f"Qo'shilgan savollar: "
            f"<b>{len(valid)}</b>\n"
            "Fan va mavzular avtomatik bog'landi.",
            reply_markup=admin_menu(),
        )

    except Exception as e:
        await m.answer(
            "<b>Import amalga oshmadi.</b>\n\n"
            f"Sabab: <code>{str(e)[:500]}</code>",
            reply_markup=cancel_kb(),
        )

    finally:
        try:
            path.unlink(
                missing_ok=True
            )
        except Exception:
            pass


@router.message(AdminState.import_file)
async def import_wrong_type(
    m: Message,
    state: FSMContext,
):
    if not is_admin(m.from_user.id):
        return

    await m.answer(
        "Excel/CSV faylini hujjat sifatida "
        "yuboring: <b>.xlsx</b> yoki <b>.csv</b>.",
        reply_markup=cancel_kb(),
    )


'''

p.write_text(
    s[:start] + new_import + s[end:],
    encoding="utf-8",
)

print("NEW IMPORT BLOCK INSTALLED")
