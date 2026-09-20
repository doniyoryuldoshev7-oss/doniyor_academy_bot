from pathlib import Path
import re
from datetime import datetime

BASE = Path("app")
MODEL = BASE / "models.py"
DB = BASE / "db.py"
ADMIN = BASE / "handlers" / "admin.py"
STUDENT = BASE / "handlers" / "student.py"

stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
backup_dir = Path("BACKUP_WORKING") / f"before_image_patch_{stamp}"
backup_dir.mkdir(parents=True, exist_ok=True)

for p in (MODEL, DB, ADMIN, STUDENT):
    (backup_dir / p.name).write_bytes(p.read_bytes())

def write(p, text):
    p.write_text(text, encoding="utf-8")

def replace_once(text, old, new, label):
    n = text.count(old)
    if n != 1:
        raise RuntimeError(f"{label}: {n} ta moslik topildi, 1 ta kutilgan")
    return text.replace(old, new, 1)

def replace_func(text, name, new_block, pattern=None):
    if pattern is None:
        pattern = rf'(?ms)(?:^@[^\n]+\n)+async def {re.escape(name)}\([^\n]*\):.*?(?=^@|\Z)'
    m = re.search(pattern, text)
    if not m:
        raise RuntimeError(f"Funksiya topilmadi: {name}")
    return text[:m.start()] + new_block.rstrip() + "\n\n" + text[m.end():]

# =========================================================
# 1. MODELS.PY
# =========================================================
model = MODEL.read_text(encoding="utf-8-sig")

if "LargeBinary" not in model:
    model = model.replace(
        "from sqlalchemy import String, Boolean, ForeignKey, DateTime, Integer, Text, BigInteger",
        "from sqlalchemy import String, Boolean, ForeignKey, DateTime, Integer, Text, BigInteger, LargeBinary",
        1
    )

if "image_data:" not in model:
    marker = '''    image_path: Mapped[str | None] = mapped_column(
        String(500),
        nullable=True
    )
'''
    if marker not in model:
        raise RuntimeError("models.py: image_path bloki topilmadi")

    model = model.replace(
        marker,
        marker + '''
    image_data: Mapped[bytes | None] = mapped_column(
        LargeBinary,
        nullable=True
    )
''',
        1
    )

write(MODEL, model)

# =========================================================
# 2. DB.PY
# =========================================================
db = DB.read_text(encoding="utf-8-sig")

if "ADD COLUMN IF NOT EXISTS image_data BYTEA" not in db:
    old_pg = """                ADD COLUMN IF NOT EXISTS correct_answer TEXT
"""
    new_pg = """                ADD COLUMN IF NOT EXISTS correct_answer TEXT,
                ADD COLUMN IF NOT EXISTS image_data BYTEA
"""
    if old_pg not in db:
        raise RuntimeError("db.py: PostgreSQL migration bloki topilmadi")
    db = db.replace(old_pg, new_pg, 1)

if '"image_data" not in existing' not in db:
    old_sqlite = '''            if "correct_answer" not in existing:
                await conn.execute(
                    text("ALTER TABLE questions ADD COLUMN correct_answer TEXT")
                )
'''
    new_sqlite = old_sqlite + '''
            if "image_data" not in existing:
                await conn.execute(
                    text("ALTER TABLE questions ADD COLUMN image_data BLOB")
                )
'''
    if old_sqlite not in db:
        raise RuntimeError("db.py: SQLite migration bloki topilmadi")
    db = db.replace(old_sqlite, new_sqlite, 1)

write(DB, db)

# =========================================================
# 3. ADMIN.PY IMPORTLAR
# =========================================================
admin = ADMIN.read_text(encoding="utf-8-sig")

if "from io import BytesIO" not in admin:
    admin = admin.replace(
        "from aiogram.fsm.context import FSMContext",
        "from aiogram.fsm.context import FSMContext\n"
        "from aiogram.types import BufferedInputFile, InputMediaPhoto\n"
        "from io import BytesIO\n"
        "import html",
        1
    )

if "async def photo_bytes(message: Message):" not in admin:
    marker = "\n\n# -------------------- SUBJECT --------------------"
    if marker not in admin:
        raise RuntimeError("admin.py: SUBJECT marker topilmadi")

    helper = '''
    
async def photo_bytes(message: Message):
    if not message.photo:
        return None

    buf = BytesIO()
    file = await message.bot.get_file(message.photo[-1].file_id)
    await message.bot.download(file, destination=buf)
    return buf.getvalue()
'''
    admin = admin.replace(marker, helper + marker, 1)

# =========================================================
# 4. ADMIN: SAVOL QO'SHISH
# =========================================================
new_q_text = '''@router.message(AdminState.add_question_text)
async def q_text(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    image_data = await photo_bytes(m)

    if image_data:
        await state.update_data(
            text=None,
            image_data=image_data,
            question_mode="image"
        )
    else:
        text = (m.text or "").strip()

        if not text:
            await m.answer("❗ Savol sifatida matn yoki surat yuboring.")
            return

        await state.update_data(
            text=text,
            image_data=None,
            question_mode="closed"
        )

    await state.set_state(AdminState.add_question_a)

    await m.answer(
        "2️⃣ <b>A variant</b>ni yuboring:",
        reply_markup=cancel_kb()
    )'''

admin = replace_func(admin, "q_text", new_q_text)

# =========================================================
# 5. ADMIN: SAVOLNI SAQLASH
# =========================================================
new_q_explanation = '''@router.message(AdminState.add_question_explanation)
async def q_explanation(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    explanation = (m.text or "").strip()
    data = await state.get_data()

    explanation = None if explanation == "-" else explanation

    async with SessionLocal() as s:
        topic = await s.get(Topic, data["topic_id"])

        if not topic:
            await state.clear()
            await m.answer(
                "❌ Mavzu topilmadi.",
                reply_markup=admin_menu()
            )
            return

        q = Question(
            topic_id=topic.id,
            text=data.get("text"),
            image_data=data.get("image_data"),
            question_mode=data.get("question_mode", "closed"),
            option_a=data["option_a"],
            option_b=data["option_b"],
            option_c=data["option_c"],
            option_d=data["option_d"],
            correct_option=data["correct_option"],
            explanation=explanation,
        )

        s.add(q)
        await s.commit()

        qid = q.id

    await state.clear()

    await m.answer(
        f"✅ <b>Savol muvaffaqiyatli qo‘shildi!</b>\\n\\n"
        f"🆔 ID: <code>{qid}</code>\\n"
        f"📖 Mavzu: <b>{topic.name}</b>\\n"
        f"✅ To‘g‘ri javob: <b>{data['correct_option']}</b>",
        reply_markup=admin_menu(),
    )'''

admin = replace_func(admin, "q_explanation", new_q_explanation)

# =========================================================
# 6. ADMIN: EXCEL IMPORT
# =========================================================
new_import = '''@router.message(AdminState.import_file, F.document)
async def import_document(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    doc = m.document
    filename = (doc.file_name or "").lower()

    if not (filename.endswith(".xlsx") or filename.endswith(".csv")):
        await m.answer(
            "❗ Faqat .xlsx yoki .csv fayl yuboring.",
            reply_markup=cancel_kb()
        )
        return

    import tempfile
    from pathlib import Path

    path = (
        Path(tempfile.gettempdir())
        / f"doniyor_import_{m.from_user.id}_{doc.file_unique_id}{Path(filename).suffix}"
    )

    try:
        await m.bot.download(doc, destination=path)

        image_by_row = {}

        # -------------------------------------------------
        # XLSX
        # -------------------------------------------------
        if filename.endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(
                path,
                read_only=False,
                data_only=True
            )

            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))

            for img in getattr(ws, "_images", []):
                try:
                    row_no = int(img.anchor._from.row) + 1
                    col_no = int(img.anchor._from.col)

                    data = img._data()

                    current = image_by_row.get(row_no)

                    # C ustun (Savol) ustuvor
                    if current is None or col_no == 2:
                        image_by_row[row_no] = data

                except Exception:
                    continue

            wb.close()

        # -------------------------------------------------
        # CSV
        # -------------------------------------------------
        else:
            import csv
            import io

            raw = path.read_bytes()
            text = None

            for enc in ("utf-8-sig", "utf-8", "cp1251"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    pass

            if text is None:
                raise ValueError("CSV kodlashini o‘qib bo‘lmadi.")

            rows = list(csv.reader(io.StringIO(text)))

        if len(rows) < 2:
            raise ValueError("Faylda savollar mavjud emas.")

        # -------------------------------------------------
        # HEADER NORMALIZATSIYA
        # -------------------------------------------------
        def normalize_header(value):
            if value is None:
                return ""

            value = str(value).strip().lower()

            for code in (
                0x2018,
                0x2019,
                0x02BB,
                0x02BC,
                0x0060
            ):
                value = value.replace(chr(code), "'")

            value = value.replace(chr(0xFEFF), "")
            value = " ".join(value.split())

            return value

        headers = [
            normalize_header(x)
            for x in rows[0]
        ]

        aliases = {
            "fan": {
                "fan",
                "subject"
            },
            "mavzu": {
                "mavzu",
                "topic"
            },
            "savol": {
                "savol",
                "question",
                "text"
            },
            "a": {
                "a",
                "variant a"
            },
            "b": {
                "b",
                "variant b"
            },
            "c": {
                "c",
                "variant c"
            },
            "d": {
                "d",
                "variant d"
            },
            "correct": {
                "to'g'ri javob",
                "togri javob",
                "to'gri javob",
                "correct",
                "correct option",
                "correct answer",
                "javob"
            },
            "izoh": {
                "izoh",
                "explanation",
                "comment"
            }
        }

        idx = {}

        for key, names in aliases.items():
            normalized_names = {
                normalize_header(x)
                for x in names
            }

            for i, h in enumerate(headers):
                if h in normalized_names:
                    idx[key] = i
                    break

        required = [
            "fan",
            "mavzu",
            "savol",
            "a",
            "b",
            "c",
            "d",
            "correct"
        ]

        missing = [
            x for x in required
            if x not in idx
        ]

        if missing:
            raise ValueError(
                "Majburiy ustunlar topilmadi: "
                + ", ".join(missing)
            )

        def cell(row, key):
            i = idx.get(key)

            if i is None:
                return ""

            if i >= len(row):
                return ""

            if row[i] is None:
                return ""

            return str(row[i]).strip()

        # -------------------------------------------------
        # SAVOLLARNI TEKSHIRISH
        # -------------------------------------------------
        valid = []
        errors = []

        for row_no, row in enumerate(
            rows[1:],
            start=2
        ):
            if not any(
                x is not None and str(x).strip()
                for x in row
            ):
                continue

            subject_name = cell(row, "fan")
            topic_name = cell(row, "mavzu")
            question = cell(row, "savol")

            a = cell(row, "a")
            b = cell(row, "b")
            copt = cell(row, "c")
            d = cell(row, "d")

            correct = cell(
                row,
                "correct"
            ).upper()

            explanation = (
                cell(row, "izoh")
                or None
            )

            image_data = (
                image_by_row.get(row_no)
                if filename.endswith(".xlsx")
                else None
            )

            if (
                not subject_name
                or not topic_name
                or not a
                or not b
                or not copt
                or not d
                or not correct
            ):
                errors.append(
                    f"{row_no}-qator: majburiy maydon bo‘sh."
                )
                continue

            # Savol matni bo'lmasa, rasm bo'lishi mumkin
            if not question and not image_data:
                errors.append(
                    f"{row_no}-qator: Savol matni yoki Savol "
                    f"ustuniga joylangan rasm kerak."
                )
                continue

            if correct not in {"A", "B", "C", "D"}:
                errors.append(
                    f"{row_no}-qator: to‘g‘ri javob A/B/C/D bo‘lishi kerak."
                )
                continue

            valid.append({
                "subject_name": subject_name,
                "topic_name": topic_name,
                "question": question or None,
                "image_data": image_data,
                "a": a,
                "b": b,
                "c": copt,
                "d": d,
                "correct": correct,
                "explanation": explanation,
            })

        if errors:
            preview = "\\n".join(
                errors[:10]
            )

            more = (
                f"\\n... yana {len(errors) - 10} ta xato"
                if len(errors) > 10
                else ""
            )

            await m.answer(
                f"⚠️ <b>Faylda xato bor.</b>\\n\\n"
                f"{preview}{more}\\n\\n"
                "Import bajarilmadi. Faylni tuzatib qayta yuboring.",
                reply_markup=cancel_kb(),
            )
            return

        if not valid:
            raise ValueError(
                "Import qilinadigan savol topilmadi."
            )

        # -------------------------------------------------
        # DATABASE
        # -------------------------------------------------
        async with SessionLocal() as s:
            subject_cache = {}
            topic_cache = {}

            for item in valid:
                subject_name = item["subject_name"]
                topic_name = item["topic_name"]

                skey = subject_name.casefold()

                subject = subject_cache.get(skey)

                if not subject:
                    subject = await s.scalar(
                        select(Subject).where(
                            func.lower(Subject.name)
                            == subject_name.lower()
                        )
                    )

                    if not subject:
                        subject = Subject(
                            name=subject_name
                        )

                        s.add(subject)
                        await s.flush()

                    subject_cache[skey] = subject

                tkey = (
                    subject.id,
                    topic_name.casefold()
                )

                topic = topic_cache.get(tkey)

                if not topic:
                    topic = await s.scalar(
                        select(Topic).where(
                            Topic.subject_id == subject.id,
                            func.lower(Topic.name)
                            == topic_name.lower()
                        )
                    )

                    if not topic:
                        topic = Topic(
                            subject_id=subject.id,
                            name=topic_name
                        )

                        s.add(topic)
                        await s.flush()

                    topic_cache[tkey] = topic

                s.add(
                    Question(
                        topic_id=topic.id,
                        text=item["question"],
                        image_data=item["image_data"],
                        question_mode=(
                            "image"
                            if item["image_data"]
                            else "closed"
                        ),
                        option_a=item["a"],
                        option_b=item["b"],
                        option_c=item["c"],
                        option_d=item["d"],
                        correct_option=item["correct"],
                        explanation=item["explanation"],
                    )
                )

            await s.commit()

        await state.clear()

        image_count = sum(
            1
            for x in valid
            if x["image_data"]
        )

        await m.answer(
            "🎉 <b>IMPORT MUVAFFAQIYATLI!</b>\\n\\n"
            f"📝 Qo‘shilgan savollar: <b>{len(valid)}</b>\\n"
            f"🖼 Rasmli savollar: <b>{image_count}</b>\\n"
            "📚 Fan va mavzular avtomatik bog‘landi.",
            reply_markup=admin_menu(),
        )

    except Exception as e:
        await m.answer(
            f"❌ <b>Import amalga oshmadi.</b>\\n\\n"
            f"Sabab: <code>{str(e)[:500]}</code>",
            reply_markup=cancel_kb(),
        )

    finally:
        try:
            path.unlink(missing_ok=True)
        except Exception:
            pass
'''

admin = replace_func(
    admin,
    "import_document",
    new_import,
    r'(?ms)^@router\.message\(AdminState\.import_file, F\.document\)\n'
    r'async def import_document\([^\n]*\):.*?'
    r'(?=^@router\.message\(AdminState\.import_file\)\nasync def import_wrong_type|\Z)'
)

# =========================================================
# 7. ADMIN: SAVOL TAHRIRLASH BOSHLANISHI
# =========================================================
new_edit_start = '''@router.callback_query(F.data.startswith("adm:edit_question:"))
async def edit_question_start(c: CallbackQuery, state: FSMContext):
    if not is_admin(c.from_user.id):
        return

    qid = int(c.data.split(":")[2])

    async with SessionLocal() as s:
        q = await s.get(Question, qid)

    if not q:
        await c.answer(
            "Savol topilmadi.",
            show_alert=True
        )
        return

    await state.update_data(
        question_id=qid,
        text=q.text,
        image_data=q.image_data,
        question_mode=(
            q.question_mode
            or ("image" if q.image_data else "closed")
        )
    )

    await state.set_state(
        AdminState.edit_question_text
    )

    old = (
        "🖼 <b>Hozirgi savol — RASM</b>"
        if q.image_data
        else
        f"Eski savol:\\n"
        f"<b>{html.escape(q.text or '')}</b>"
    )

    await c.message.edit_text(
        "✏️ <b>SAVOLNI TAHRIRLASH</b>\\n\\n"
        f"{old}\\n\\n"
        "1️⃣ Yangi savol sifatida "
        "<b>matn yoki surat</b> yuboring:",
        reply_markup=cancel_kb(),
    )

    await c.answer()
'''

admin = replace_func(
    admin,
    "edit_question_start",
    new_edit_start
)

# =========================================================
# 8. ADMIN: SAVOL TAHRIRLASHDA MATN / RASM
# =========================================================
new_edit_text = '''@router.message(AdminState.edit_question_text)
async def edit_question_text_handler(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    image_data = await photo_bytes(m)

    if image_data:
        await state.update_data(
            text=None,
            image_data=image_data,
            question_mode="image"
        )
    else:
        value = (m.text or "").strip()

        if not value:
            await m.answer(
                "❗ Savol sifatida matn yoki surat yuboring."
            )
            return

        await state.update_data(
            text=value,
            image_data=None,
            question_mode="closed"
        )

    await state.set_state(
        AdminState.edit_question_a
    )

    await m.answer(
        "2️⃣ <b>A variant</b>ni yuboring:",
        reply_markup=cancel_kb()
    )
'''

admin = replace_func(
    admin,
    "edit_question_text_handler",
    new_edit_text
)

# =========================================================
# 9. ADMIN: TAHRIRLASHNI SAQLASH
# =========================================================
new_edit_d = '''@router.message(AdminState.edit_question_d)
async def edit_question_d_handler(m: Message, state: FSMContext):
    if not is_admin(m.from_user.id):
        return

    value = (m.text or "").strip()

    if not value:
        await m.answer(
            "❗ Variant bo‘sh bo‘lishi mumkin emas."
        )
        return

    await state.update_data(
        option_d=value
    )

    data = await state.get_data()

    async with SessionLocal() as s:
        q = await s.get(
            Question,
            data["question_id"]
        )

        if not q:
            await state.clear()

            await m.answer(
                "❌ Savol topilmadi.",
                reply_markup=admin_menu()
            )
            return

        q.text = data.get("text")
        q.image_data = data.get("image_data")

        q.question_mode = data.get(
            "question_mode",
            "closed"
        )

        q.option_a = data["option_a"]
        q.option_b = data["option_b"]
        q.option_c = data["option_c"]
        q.option_d = data["option_d"]

        await s.commit()

        topic_id = q.topic_id

    await state.clear()

    await m.answer(
        "✅ <b>SAVOL MUVAFFAQIYATLI TAHRIRLANDI!</b>",
        reply_markup=InlineKeyboardMarkup(
            inline_keyboard=[
                [
                    InlineKeyboardButton(
                        text="⬅️ Savollar",
                        callback_data=f"adm:questions:{topic_id}"
                    )
                ],
                [
                    InlineKeyboardButton(
                        text="🏠 Admin panel",
                        callback_data="admin"
                    )
                ]
            ]
        )
    )
'''

admin = replace_func(
    admin,
    "edit_question_d_handler",
    new_edit_d
)

# =========================================================
# 10. ADMIN: SAVOLNI KO'RISH
# =========================================================
new_catalog_question = '''@router.callback_query(F.data.startswith("adm:question:"))
async def catalog_question(c: CallbackQuery):
    if not is_admin(c.from_user.id):
        return

    qid = int(c.data.split(":")[2])

    async with SessionLocal() as s:
        q = await s.get(Question, qid)

        if not q:
            await c.answer(
                "Savol topilmadi.",
                show_alert=True
            )
            return

        topic = await s.get(
            Topic,
            q.topic_id
        )

    text = (
        "📝 <b>SAVOL</b>\n"
        "━━━━━━━━━━━━━━\n\n"
        f"📖 <b>{html.escape(topic.name if topic else '')}</b>\n\n"
        +
        (
            "🖼 <b>RASMLI SAVOL</b>"
            if q.image_data
            else
            f"❓ <b>{html.escape(q.text or '')}</b>"
        )
        +
        "\n\n"
        f"A) {html.escape(q.option_a or '')}\n"
        f"B) {html.escape(q.option_b or '')}\n"
        f"C) {html.escape(q.option_c or '')}\n"
        f"D) {html.escape(q.option_d or '')}\n\n"
        f"✅ To‘g‘ri javob: <b>{q.correct_option}</b>"
    )

    if q.explanation:
        text += (
            f"\n\n💡 <b>Izoh:</b> "
            f"{html.escape(q.explanation)}"
        )

    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="✏️ Tahrirlash",
                    callback_data=f"adm:edit_question:{q.id}"
                )
            ],
            [
                InlineKeyboardButton(
                    text="🗑 O‘chirish",
                    callback_data=f"adm:delete_question:{q.id}"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⬅️ Savollar",
                    callback_data=f"adm:questions:{q.topic_id}"
                )
            ],
            [
                InlineKeyboardButton(
                    text="🏠 Admin panel",
                    callback_data="admin"
                )
            ]
        ]
    )

    if q.image_data:
        caption = text[:1024]

        media = InputMediaPhoto(
            media=BufferedInputFile(
                q.image_data,
                filename=f"question_{q.id}.jpg"
            ),
            caption=caption,
            parse_mode="HTML"
        )

        if c.message.photo:
            await c.message.edit_media(
                media=media,
                reply_markup=kb
            )
        else:
            try:
                await c.message.delete()
            except Exception:
                pass

            await c.message.answer_photo(
                BufferedInputFile(
                    q.image_data,
                    filename=f"question_{q.id}.jpg"
                ),
                caption=caption,
                reply_markup=kb,
                parse_mode="HTML"
            )

    else:
        if c.message.photo:
            try:
                await c.message.delete()
            except Exception:
                pass

            await c.message.answer(
                text,
                reply_markup=kb,
                parse_mode="HTML"
            )
        else:
            await c.message.edit_text(
                text,
                reply_markup=kb
            )

    await c.answer()
'''

admin = replace_func(
    admin,
    "catalog_question",
    new_catalog_question
)

write(ADMIN, admin)

# =========================================================
# 11. STUDENT.PY IMPORT
# =========================================================
student = STUDENT.read_text(encoding="utf-8-sig")

if "BufferedInputFile" not in student:
    student = student.replace(
        "from aiogram.types import CallbackQuery,InlineKeyboardMarkup,InlineKeyboardButton",
        "from aiogram.types import CallbackQuery,InlineKeyboardMarkup,InlineKeyboardButton,BufferedInputFile",
        1
    )

# =========================================================
# 12. STUDENT: RENDER
# =========================================================
new_render = '''async def render(call, state):
    d = await state.get_data()

    async with SessionLocal() as s:
        q = await s.get(
            Question,
            d["qids"][d["index"]]
        )

        topic = await s.get(
            Topic,
            d["topic_id"]
        )

    n = d["index"] + 1
    total = len(d["qids"])
    progress = int(n / total * 100)

    body = (
        f"📖 <b>{topic.name}</b>\\n\\n"
        f"━━━━━━━━━━━━\\n"
        f"❓ <b>Savol: {n}/{total}</b>\\n"
        f"📊 <b>Progress: {progress}%</b>\\n"
        f"━━━━━━━━━━━━\\n\\n"
    )

    body += (
        q.text or "🖼 <b>Rasmli savol</b>"
    )

    body += (
        "\\n\\n"
        f"<b>A)</b> {q.option_a or ''}\\n"
        f"<b>B)</b> {q.option_b or ''}\\n"
        f"<b>C)</b> {q.option_c or ''}\\n"
        f"<b>D)</b> {q.option_d or ''}"
    )

    kb = answer_kb(
        topic.id,
        q.id
    )

    if q.image_data:
        caption = body[:1024]

        if call.message.photo:
            await call.message.edit_caption(
                caption=caption,
                reply_markup=kb,
                parse_mode="HTML"
            )
        else:
            try:
                await call.message.delete()
            except Exception:
                pass

            await call.message.answer_photo(
                BufferedInputFile(
                    q.image_data,
                    filename=f"question_{q.id}.jpg"
                ),
                caption=caption,
                reply_markup=kb,
                parse_mode="HTML"
            )

    else:
        if call.message.photo:
            try:
                await call.message.delete()
            except Exception:
                pass

            await call.message.answer(
                body,
                reply_markup=kb,
                parse_mode="HTML"
            )
        else:
            await call.message.edit_text(
                body,
                reply_markup=kb
            )
'''

student = replace_func(
    student,
    "render",
    new_render
)

# =========================================================
# 13. STUDENT: ANSWERDA RASMNI SAQLAB QOLISH
# =========================================================
start = student.find(
    '@router.callback_query(QuizState.active, F.data.startswith("answer:"))'
)

end = student.find(
    "@router.callback_query(F.data=='quiz:next')",
    start
)

if start == -1 or end == -1:
    raise RuntimeError(
        "student.py: answer funksiyasi topilmadi"
    )

answer_block = student[start:end]

old_retry = '''            await c.message.edit_text(
                f"❌ <b>Yana bir bor urinib ko‘ring!</b>\\n\\n"
                f"❓ <b>{q.text}</b>\\n\\n"
                f"<b>A)</b> {q.option_a}\\n"
                f"<b>B)</b> {q.option_b}\\n"
                f"<b>C)</b> {q.option_c}\\n"
                f"<b>D)</b> {q.option_d}\\n\\n"
                f"💡 <i>To‘g‘ri javobni topmaguningizcha davom etamiz.</i>",
                reply_markup=keyboard
            )'''

new_retry = '''            retry_text = (
                f"❌ <b>Yana bir bor urinib ko‘ring!</b>\\n\\n"
                f"❓ <b>{q.text or '🖼 Rasmli savol'}</b>\\n\\n"
                f"<b>A)</b> {q.option_a}\\n"
                f"<b>B)</b> {q.option_b}\\n"
                f"<b>C)</b> {q.option_c}\\n"
                f"<b>D)</b> {q.option_d}\\n\\n"
                f"💡 <i>To‘g‘ri javobni topmaguningizcha davom etamiz.</i>"
            )

            if q.image_data and c.message.photo:
                await c.message.edit_caption(
                    caption=retry_text[:1024],
                    reply_markup=keyboard,
                    parse_mode="HTML"
                )

            elif q.image_data:
                try:
                    await c.message.delete()
                except Exception:
                    pass

                await c.message.answer_photo(
                    BufferedInputFile(
                        q.image_data,
                        filename=f"question_{q.id}.jpg"
                    ),
                    caption=retry_text[:1024],
                    reply_markup=keyboard,
                    parse_mode="HTML"
                )

            else:
                await c.message.edit_text(
                    retry_text,
                    reply_markup=keyboard
                )'''

answer_block = replace_once(
    answer_block,
    old_retry,
    new_retry,
    "student retry"
)

old_success = '''        await c.message.edit_text(
            f"🎉 <b>BARAKALLA!</b>\\n"
            f"━━━━━━━━━━━━━━\\n\\n"
            f"✅ Siz savolga to‘g‘ri javob berdingiz!\\n\\n"
            f"❓ <b>{q.text}</b>\\n\\n"
            f"✅ To‘g‘ri javob:\\n"
            f"<b>{q.correct_option}) {opt(q, q.correct_option)}</b>\\n\\n"
            f"🏆 <b>Bu savol o‘zlashtirildi!</b>\\n\\n"
            f"Endi u <b>“Mening xatolarim”</b> bo‘limida chiqmaydi.",
            reply_markup=InlineKeyboardMarkup(
                inline_keyboard=[
                    [
                        InlineKeyboardButton(
                            text="❌ Mening xatolarim",
                            callback_data="my_errors"
                        )
                    ],
                    [
                        InlineKeyboardButton(
                            text="🏠 Bosh menyu",
                            callback_data="home"
                        )
                    ]
                ]
            )
        )'''

new_success = '''        success_text = (
            f"🎉 <b>BARAKALLA!</b>\\n"
            f"━━━━━━━━━━━━━━\\n\\n"
            f"✅ Siz savolga to‘g‘ri javob berdingiz!\\n\\n"
            f"❓ <b>{q.text or '🖼 Rasmli savol'}</b>\\n\\n"
            f"✅ To‘g‘ri javob:\\n"
            f"<b>{q.correct_option}) {opt(q, q.correct_option)}</b>\\n\\n"
            f"🏆 <b>Bu savol o‘zlashtirildi!</b>\\n\\n"
            f"Endi u <b>“Mening xatolarim”</b> bo‘limida chiqmaydi."
        )

        success_kb = InlineKeyboardMarkup(
            inline_keyboard=[
                [
                    InlineKeyboardButton(
                        text="❌ Mening xatolarim",
                        callback_data="my_errors"
                    )
                ],
                [
                    InlineKeyboardButton(
                        text="🏠 Bosh menyu",
                        callback_data="home"
                    )
                ]
            ]
        )

        if q.image_data and c.message.photo:
            await c.message.edit_caption(
                caption=success_text[:1024],
                reply_markup=success_kb,
                parse_mode="HTML"
            )

        elif q.image_data:
            try:
                await c.message.delete()
            except Exception:
                pass

            await c.message.answer_photo(
                BufferedInputFile(
                    q.image_data,
                    filename=f"question_{q.id}.jpg"
                ),
                caption=success_text[:1024],
                reply_markup=success_kb,
                parse_mode="HTML"
            )

        else:
            await c.message.edit_text(
                success_text,
                reply_markup=success_kb
            )'''

answer_block = replace_once(
    answer_block,
    old_success,
    new_success,
    "student retry success"
)

old_result = '''    await c.message.edit_text(
        f"{result}\\n\\n"
        f"📊 Hozirgi natija: "
        f"<b>{correct}/{answered}</b>",
        reply_markup=InlineKeyboardMarkup(
            inline_keyboard=buttons
        )
    )'''

new_result = '''    result_text = (
        f"{result}\\n\\n"
        f"📊 Hozirgi natija: "
        f"<b>{correct}/{answered}</b>"
    )

    result_kb = InlineKeyboardMarkup(
        inline_keyboard=buttons
    )

    if q.image_data and c.message.photo:
        await c.message.edit_caption(
            caption=result_text[:1024],
            reply_markup=result_kb,
            parse_mode="HTML"
        )

    elif q.image_data:
        try:
            await c.message.delete()
        except Exception:
            pass

        await c.message.answer_photo(
            BufferedInputFile(
                q.image_data,
                filename=f"question_{q.id}.jpg"
            ),
            caption=result_text[:1024],
            reply_markup=result_kb,
            parse_mode="HTML"
        )

    else:
        await c.message.edit_text(
            result_text,
            reply_markup=result_kb
        )'''

answer_block = replace_once(
    answer_block,
    old_result,
    new_result,
    "student normal result"
)

student = (
    student[:start]
    + answer_block
    + student[end:]
)

write(STUDENT, student)

print()
print("======================================")
print("IMAGE PATCH MUVAFFAQIYATLI QO'LLANDI")
print("======================================")
print("BACKUP:")
print(backup_dir)
print()

for p in (MODEL, DB, ADMIN, STUDENT):
    print(f"{p}: {p.stat().st_size} bytes")

